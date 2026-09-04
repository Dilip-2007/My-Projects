import streamlit as st
from google import genai
from google.genai import types
from dotenv import load_dotenv
from pptx import Presentation
from pptx.util import Inches, Pt
import os
import json
import io


# =========================================================
# LOAD ENVIRONMENT VARIABLES
# =========================================================

load_dotenv()

api_key = os.getenv("GEMINI_API_KEY")


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="AI Executive Slide Generator",
    page_icon="📊",
    layout="wide"
)


# =========================================================
# CHECK API KEY
# =========================================================

if not api_key:

    st.error(
        "Gemini API key is missing. "
        "Please add GEMINI_API_KEY to your .env file."
    )

    st.stop()


# =========================================================
# GEMINI CLIENT
# =========================================================

client = genai.Client(api_key=api_key)


# =========================================================
# TITLE
# =========================================================

st.title("📊 Autonomous Research & Executive Slide Deck Generator")

st.write(
    "Enter an industry or business topic. Gemini will research it, "
    "create a 5-slide executive presentation, and generate a PowerPoint file."
)


# =========================================================
# TOPIC INPUT
# =========================================================

topic = st.text_input(
    "🔎 Enter your research topic",
    placeholder="Example: Artificial Intelligence in Healthcare"
)


# =========================================================
# GENERATE BUTTON
# =========================================================

if st.button("🚀 Research & Generate Presentation"):

    if topic.strip() == "":
        st.warning("Please enter a topic.")

    else:

        prompt = f"""
You are an autonomous business research analyst.

Research the following topic using current and reliable web information:

TOPIC:
{topic}

Create an executive-level presentation containing EXACTLY 5 slides.

The slides must be:

1. Title & Executive Overview
2. Current Market / Industry Situation
3. Key Trends and Developments
4. Opportunities and Challenges
5. Future Outlook & Key Takeaways

For every slide provide:

- slide_number
- title
- 3 to 5 concise bullet points
- a short speaker_note
- source URLs used for the information

Requirements:

- Use current information.
- Prefer reliable sources.
- Do not invent statistics.
- Keep the content concise and suitable for an executive presentation.
- Make the bullets understandable without additional explanation.

Return ONLY valid JSON.

JSON structure:

{{
    "topic": "{topic}",
    "slides": [
        {{
            "slide_number": 1,
            "title": "...",
            "bullets": [
                "...",
                "...",
                "..."
            ],
            "speaker_note": "...",
            "sources": [
                {{
                    "title": "...",
                    "url": "..."
                }}
            ]
        }}
    ]
}}
"""


        # =================================================
        # GEMINI RESEARCH
        # =================================================

        try:

            with st.spinner(
                "🔎 Gemini is researching the topic and creating your presentation..."
            ):

                response = client.models.generate_content(
                    model="gemini-3.6-flash",
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        tools=[
                            types.Tool(
                                google_search=types.GoogleSearch()
                            )
                        ],
                        response_mime_type="application/json"
                    )
                )


            # =================================================
            # PARSE JSON
            # =================================================

            data = json.loads(response.text)

            slides = data["slides"]


            # =================================================
            # VALIDATE 5 SLIDES
            # =================================================

            if len(slides) != 5:

                st.error(
                    "Gemini did not generate exactly 5 slides."
                )

                st.stop()


            # =================================================
            # SHOW SUCCESS
            # =================================================

            st.success(
                "✅ Research completed and 5-slide presentation generated!"
            )


            # =================================================
            # SLIDE PREVIEW
            # =================================================

            st.header("👀 Presentation Preview")


            for slide in slides:

                with st.container(border=True):

                    st.subheader(
                        f"Slide {slide['slide_number']}: "
                        f"{slide['title']}"
                    )

                    for bullet in slide["bullets"]:

                        st.write(
                            f"• {bullet}"
                        )

                    st.caption(
                        f"🗣️ Speaker note: "
                        f"{slide['speaker_note']}"
                    )


                    if slide.get("sources"):

                        st.write("🔗 Sources:")

                        for source in slide["sources"]:

                            st.write(
                                f"- {source['title']}: "
                                f"{source['url']}"
                            )


            # =================================================
            # CREATE POWERPOINT
            # =================================================

            st.header("📥 Download Presentation")


            prs = Presentation()


            # =================================================
            # CREATE TITLE SLIDE
            # =================================================

            title_slide = prs.slides.add_slide(
                prs.slide_layouts[0]
            )

            title_slide.shapes.title.text = topic

            title_slide.placeholders[1].text = (
                "Autonomous Research & Executive Briefing"
            )


            # =================================================
            # CREATE CONTENT SLIDES
            # =================================================

            for slide_data in slides[1:]:

                slide = prs.slides.add_slide(
                    prs.slide_layouts[1]
                )

                slide.shapes.title.text = (
                    slide_data["title"]
                )


                # ---------------------------------------------
                # BULLET CONTENT
                # ---------------------------------------------

                text_frame = (
                    slide.placeholders[1].text_frame
                )

                text_frame.clear()


                for index, bullet in enumerate(
                    slide_data["bullets"]
                ):

                    if index == 0:

                        paragraph = text_frame.paragraphs[0]

                    else:

                        paragraph = text_frame.add_paragraph()


                    paragraph.text = bullet

                    paragraph.level = 0

                    paragraph.font.size = Pt(22)


                # ---------------------------------------------
                # SOURCES
                # ---------------------------------------------

                sources = slide_data.get(
                    "sources",
                    []
                )


                if sources:

                    source_text = "\n\nSources:\n"

                    for source in sources:

                        source_text += (
                            f"{source['title']}: "
                            f"{source['url']}\n"
                        )


                    textbox = slide.shapes.add_textbox(
                        Inches(0.5),
                        Inches(6.3),
                        Inches(9),
                        Inches(0.7)
                    )


                    tf = textbox.text_frame

                    tf.text = source_text

                    for paragraph in tf.paragraphs:

                        paragraph.font.size = Pt(8)


            # =================================================
            # SAVE PPTX TO MEMORY
            # =================================================

            pptx_buffer = io.BytesIO()

            prs.save(pptx_buffer)

            pptx_buffer.seek(0)


            # =================================================
            # DOWNLOAD BUTTON
            # =================================================

            st.download_button(
                label="⬇️ Download Executive Presentation",
                data=pptx_buffer,
                file_name="executive_research_deck.pptx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "presentationml.presentation"
                )
            )


        # =====================================================
        # ERROR HANDLING
        # =====================================================

        except json.JSONDecodeError:

            st.error(
                "Gemini returned an invalid JSON response. "
                "Please try again."
            )

        except Exception as e:

            st.error(
                f"Something went wrong: {e}"
            )